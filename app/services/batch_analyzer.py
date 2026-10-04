import csv
from datetime import date, datetime, time
from io import BytesIO, StringIO
from itertools import islice
import re
from typing import Optional

from fastapi import HTTPException, UploadFile
from openpyxl import load_workbook
import xlrd

from app.services.llm_analyzer import (
    SPREADSHEET_FIELDS,
    OllamaUnavailableError,
    match_spreadsheet_columns,
)
from app.services.readability import (
    field_length_error,
    identical_fields_error,
    normalize_text,
    readability_error,
)


FIELD_ALIASES = {
    "title": {"title", "bug title", "bug name"},
    "description": {"description", "bug description"},
    "steps_to_reproduce": {
        "steps to reproduce",
        "steps",
        "reproduction steps",
        "steps to reproduce the bug",
    },
    "expected_result": {
        "expected result",
        "expected behavior",
        "expected behaviour",
    },
    "actual_result": {
        "actual result",
        "actual behavior",
        "actual behaviour",
    },
    "severity": {
        "severity",
        "bug severity",
        "issue severity",
    },
    "priority": {
        "priority",
        "bug priority",
        "issue priority",
    },
    "category": {
        "category",
        "bug category",
        "issue category",
        "type",
        "bug type",
    },
}


REQUIRED_FIELDS = (
    "title",
    "description",
    "steps_to_reproduce",
    "expected_result",
    "actual_result",
)


HEADER_SEARCH_ROWS = 20

# What the LLM sees when the headers are not recognized: the header row and a
# few shortened sample rows.
COLUMN_SAMPLE_ROWS = 3
COLUMN_SAMPLE_CHARS = 80

# Upload limits that keep a single batch from exhausting memory or keeping
# the LLM busy for hours.
MAX_EXCEL_BYTES = 5 * 1024 * 1024
MAX_SHEET_ROWS = 2000
MAX_BUG_RECORDS = 500


def _format_number(value: float) -> str:
    if float(value).is_integer():
        return str(int(value))
    return format(value, ".15g")


def _format_datetime(value: datetime) -> str:
    if value.time() == time():
        return value.date().isoformat()
    timespec = "minutes" if value.second == 0 else "seconds"
    return value.isoformat(sep=" ", timespec=timespec)


def _format_cell(value: object, number_format: str = "General") -> object:
    """Render numbers, percentages and dates the way Excel displays them."""
    if value is None or isinstance(value, (bool, str)):
        return value
    if isinstance(value, datetime):
        return _format_datetime(value)
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, (int, float)):
        if "%" in number_format:
            return f"{_format_number(value * 100)}%"
        return _format_number(value)
    return value


def _unsaved_formula_cells(
    content: bytes, rows: list[tuple[object, ...]]
) -> frozenset[tuple[int, int]]:
    """Find formula cells whose result was never saved.

    Excel stores each formula's result when it saves a file; files written by
    scripts often contain only the formula, which reads as an empty cell.
    """
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=False)
    try:
        cells = set()
        for row_number, row in enumerate(
            islice(workbook.active.iter_rows(min_row=1, values_only=True), len(rows)),
            start=1,
        ):
            saved = rows[row_number - 1]
            for index, value in enumerate(row):
                is_formula = isinstance(value, str) and value.startswith("=")
                no_result = index >= len(saved) or saved[index] is None
                if is_formula and no_result:
                    cells.add((row_number, index))
        return frozenset(cells)
    finally:
        workbook.close()


# Encrypted Office files are OLE containers, not the zip of a normal .xlsx.
OLE_SIGNATURE = b"\xd0\xcf\x11\xe0"

PASSWORD_PROTECTED_ERROR = (
    "The Excel file is password-protected. Remove the password, save the "
    "file and upload it again."
)


def _csv_rows(content: bytes) -> list[tuple[object, ...]]:
    """Read a CSV export, guessing the encoding and the delimiter."""
    for encoding in ("utf-8-sig", "cp1254"):
        try:
            text = content.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        text = content.decode("latin-1")

    try:
        dialect = csv.Sniffer().sniff(text[:5000], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel

    rows = []
    for row in csv.reader(StringIO(text), dialect):
        rows.append(tuple(row))
        if len(rows) > MAX_SHEET_ROWS:
            break
    return rows


def _xls_cell_value(workbook: xlrd.book.Book, cell: xlrd.sheet.Cell) -> object:
    if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK, xlrd.XL_CELL_ERROR):
        return None
    if cell.ctype == xlrd.XL_CELL_BOOLEAN:
        return bool(cell.value)
    if cell.ctype == xlrd.XL_CELL_DATE:
        return _format_cell(xlrd.xldate.xldate_as_datetime(cell.value, workbook.datemode))
    if cell.ctype == xlrd.XL_CELL_NUMBER:
        format_key = workbook.xf_list[cell.xf_index].format_key
        return _format_cell(cell.value, workbook.format_map[format_key].format_str)
    return cell.value


def _normalize_header(value: object) -> str:
    return re.sub(r"[^a-z0-9]+", " ", str(value or "").strip().lower()).strip()


def _header_indexes(row: tuple[object, ...]) -> dict[str, int]:
    headers = [_normalize_header(value) for value in row]
    indexes: dict[str, int] = {}
    for field, aliases in FIELD_ALIASES.items():
        match = next((i for i, header in enumerate(headers) if header in aliases), None)
        if match is not None:
            indexes[field] = match
    return indexes


def _find_header(rows: list[tuple[object, ...]]) -> tuple[int, dict[str, int]]:
    """Find the header row, allowing report titles or notes above it."""
    best_row_index = 0
    best_indexes: dict[str, int] = {}
    best_score = 0
    for row_index, row in enumerate(rows[:HEADER_SEARCH_ROWS]):
        indexes = _header_indexes(row)
        score = sum(field in indexes for field in REQUIRED_FIELDS)
        if score > best_score:
            best_row_index, best_indexes, best_score = row_index, indexes, score
        if score == len(REQUIRED_FIELDS):
            break
    return best_row_index, best_indexes


def _cell_text(value: object) -> str:
    return "" if value is None else str(value).strip()


def _candidate_header_row(rows: list[tuple[object, ...]]) -> Optional[int]:
    """First row with at least as many filled cells as required fields."""
    for row_index, row in enumerate(rows[:HEADER_SEARCH_ROWS]):
        if sum(1 for value in row if _cell_text(value)) >= len(REQUIRED_FIELDS):
            return row_index
    return None


async def _match_columns_with_llm(
    rows: list[tuple[object, ...]],
) -> Optional[tuple[int, dict[str, int]]]:
    """Let the LLM match headers it does not know, for example other languages.

    The answer is only used when every required field gets its own existing
    column; optional fields that clash with another field are dropped.
    """
    header_index = _candidate_header_row(rows)
    if header_index is None:
        return None

    headers = [_cell_text(value) for value in rows[header_index]]
    samples = [
        [_cell_text(value)[:COLUMN_SAMPLE_CHARS] for value in row]
        for row in rows[header_index + 1 :]
        if any(_cell_text(value) for value in row)
    ][:COLUMN_SAMPLE_ROWS]

    mapping = await match_spreadsheet_columns(headers, samples)

    def is_column(index: object) -> bool:
        return isinstance(index, int) and 0 <= index < len(headers) and headers[index]

    required = {field: mapping.get(field) for field in REQUIRED_FIELDS}
    if not all(is_column(index) for index in required.values()):
        return None
    if len(set(required.values())) != len(required):
        return None

    indexes = dict(required)
    for field in SPREADSHEET_FIELDS:
        index = mapping.get(field)
        if field not in indexes and is_column(index) and index not in indexes.values():
            indexes[field] = index

    return header_index, indexes


def build_record(row_number: int, values: dict[str, str]) -> dict:
    """Validate one bug's values; rows with an error are not sent to the LLM."""
    missing_values = [
        field.replace("_", " ")
        for field in REQUIRED_FIELDS
        if not values.get(field, "").strip()
    ]
    if missing_values:
        error = f"Missing required values: {', '.join(missing_values)}."
    else:
        error = (
            field_length_error(values)
            or readability_error(values)
            or identical_fields_error(values)
        )
    return {"row": row_number, "values": values, "error": error, "duplicate_of": None}


def _unsaved_formula_error(fields: list[str]) -> str:
    labels = ", ".join(field.replace("_", " ") for field in fields)
    return (
        f"Formula without a saved result in: {labels}. Open the file in Excel "
        "and save it again so the formula results are stored."
    )


async def _records_from_rows(
    rows: list[tuple[object, ...]],
    unsaved_formulas: frozenset[tuple[int, int]] = frozenset(),
) -> tuple[list[dict], Optional[dict[str, str]]]:
    """Return the bug records and, when the LLM matched the columns, which
    header was used for each field so the user can check it.

    unsaved_formulas holds (Excel row, column index) of formula cells whose
    result was never saved; they read as empty, so the row gets an
    explanation instead of "missing values".
    """
    if not rows:
        raise HTTPException(status_code=422, detail="The Excel file is empty.")

    header_index, indexes = _find_header(rows)
    detected_columns = None

    missing = [field for field in REQUIRED_FIELDS if field not in indexes]
    if missing:
        hint = ""
        try:
            matched = await _match_columns_with_llm(rows)
        except OllamaUnavailableError:
            matched = None
            hint = (
                " Headers that are not in English are matched by the LLM, "
                "but Ollama is unavailable."
            )
        except RuntimeError:
            matched = None

        if matched is None:
            labels = ", ".join(field.replace("_", " ") for field in missing)
            raise HTTPException(
                status_code=422,
                detail=f"Missing required Excel columns: {labels}.{hint}",
            )

        header_index, indexes = matched
        headers = rows[header_index]
        detected_columns = {
            field: _cell_text(headers[index]) for field, index in indexes.items()
        }

    records = []
    for row_number, row in enumerate(
        rows[header_index + 1 :], start=header_index + 2
    ):
        values = {
            field: str(row[index]).strip()
            if index < len(row) and row[index] is not None
            else ""
            for field, index in indexes.items()
        }
        formula_fields = [
            field
            for field in REQUIRED_FIELDS
            if (row_number, indexes[field]) in unsaved_formulas
        ]
        if not any(values.values()) and not formula_fields:
            continue

        record = build_record(row_number, values)
        if formula_fields:
            record["error"] = _unsaved_formula_error(formula_fields)
        records.append(record)

    if not records:
        raise HTTPException(
            status_code=422, detail="The Excel file contains no bug records."
        )

    if len(records) > MAX_BUG_RECORDS:
        raise HTTPException(
            status_code=422,
            detail=(
                f"The Excel file contains {len(records)} bug records; "
                f"the maximum is {MAX_BUG_RECORDS}."
            ),
        )

    _mark_duplicates(records)
    return records, detected_columns


def _mark_duplicates(records: list[dict]) -> None:
    """Point repeated bugs at the 1-based number of their first occurrence."""
    first_seen: dict[tuple[str, ...], int] = {}
    for bug_number, record in enumerate(records, start=1):
        if record["error"]:
            continue
        key = tuple(
            normalize_text(record["values"][field]) for field in REQUIRED_FIELDS
        )
        if key in first_seen:
            record["duplicate_of"] = first_seen[key]
        else:
            first_seen[key] = bug_number


async def parse_bug_spreadsheet(
    file: UploadFile,
) -> tuple[list[dict], Optional[dict[str, str]]]:
    filename = (file.filename or "").lower()
    content = await file.read(MAX_EXCEL_BYTES + 1)
    if len(content) > MAX_EXCEL_BYTES:
        raise HTTPException(
            status_code=413,
            detail=f"The Excel file is larger than {MAX_EXCEL_BYTES // (1024 * 1024)} MB.",
        )

    if not content:
        raise HTTPException(status_code=422, detail="The uploaded file is empty.")

    if filename.endswith(".xlsx") and content.startswith(OLE_SIGNATURE):
        raise HTTPException(status_code=422, detail=PASSWORD_PROTECTED_ERROR)

    too_many_rows = HTTPException(
        status_code=422,
        detail=f"The Excel sheet has more than {MAX_SHEET_ROWS} rows.",
    )
    try:
        if filename.endswith(".xlsx"):
            workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
            try:
                rows = [
                    tuple(
                        _format_cell(cell.value, cell.number_format or "General")
                        for cell in row
                    )
                    for row in islice(
                        workbook.active.iter_rows(min_row=1), MAX_SHEET_ROWS + 1
                    )
                ]
            finally:
                workbook.close()
            if len(rows) > MAX_SHEET_ROWS:
                raise too_many_rows
            unsaved_formulas = _unsaved_formula_cells(content, rows)
        elif filename.endswith(".xls"):
            workbook = xlrd.open_workbook(file_contents=content, formatting_info=True)
            sheet = workbook.sheet_by_index(0)
            if sheet.nrows > MAX_SHEET_ROWS:
                raise too_many_rows
            rows = [
                tuple(_xls_cell_value(workbook, cell) for cell in sheet.row(index))
                for index in range(sheet.nrows)
            ]
            # xlrd only exposes saved formula results, not the formulas.
            unsaved_formulas = frozenset()
        elif filename.endswith(".csv"):
            rows = _csv_rows(content)
            if len(rows) > MAX_SHEET_ROWS:
                raise too_many_rows
            unsaved_formulas = frozenset()
        else:
            raise HTTPException(
                status_code=415,
                detail="Only .xlsx, .xls and .csv files are supported for batch analysis.",
            )
    except HTTPException:
        raise
    except xlrd.biffh.XLRDError as exc:
        if "encrypted" in str(exc).lower():
            raise HTTPException(
                status_code=422, detail=PASSWORD_PROTECTED_ERROR
            ) from exc
        raise HTTPException(
            status_code=422, detail="The uploaded Excel file could not be read."
        ) from exc
    except Exception as exc:
        raise HTTPException(
            status_code=422, detail="The uploaded Excel file could not be read."
        ) from exc

    return await _records_from_rows(rows, unsaved_formulas)
