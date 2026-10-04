"""Excel reports of single bug and batch analysis results."""

from collections import Counter
from datetime import datetime
from io import BytesIO
from typing import Iterable, Optional

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.schemas.analysis import BugAnalysis
from app.services.steps import split_steps


HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="2F6DF6")
SECTION_FONT = Font(bold=True, size=12)
WRAP = Alignment(wrap_text=True, vertical="top")

STATUS_LABELS = {
    "analyzed": "Analyzed",
    "failed": "Failed",
    "duplicate": "Duplicate",
    "not_analyzed": "Not analyzed",
}

SCENARIO_COLUMNS = (
    ("ID", 10),
    ("Category", 22),
    ("Type", 12),
    ("Scenario", 50),
    ("Expected Result", 40),
    ("Purpose", 40),
    ("Priority", 10),
)


def _text(value: object) -> object:
    """Cell value: numbers stay numbers, text loses characters Excel rejects."""
    if value is None or isinstance(value, (int, float)):
        return value
    return ILLEGAL_CHARACTERS_RE.sub("", str(value))


def _set(sheet: Worksheet, row: int, column: int, value: object):
    cell = sheet.cell(row=row, column=column, value=_text(value))
    # Report text comes from users and the LLM; text starting with "=" must
    # stay text instead of becoming a formula that Excel runs.
    if isinstance(cell.value, str):
        cell.data_type = "s"
    cell.alignment = WRAP
    return cell


def _numbered(text: str) -> str:
    return "\n".join(
        f"{number}. {step}" for number, step in enumerate(split_steps(text), start=1)
    )


def _lines(items: Iterable[str]) -> str:
    return "\n".join(f"• {item}" for item in items)


def _table(
    sheet: Worksheet,
    columns: Iterable[tuple[str, int]],
    rows: Iterable[Iterable[object]],
) -> None:
    """Header row, data rows, filter and frozen header."""
    columns = list(columns)
    for index, (title, width) in enumerate(columns, start=1):
        cell = _set(sheet, 1, index, title)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        sheet.column_dimensions[get_column_letter(index)].width = width

    last_row = 1
    for last_row, values in enumerate(rows, start=2):
        for index, value in enumerate(values, start=1):
            _set(sheet, last_row, index, value)

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{last_row}"


def _confidence(sheet: Worksheet, column: int, first_row: int = 2) -> None:
    for row in range(first_row, sheet.max_row + 1):
        sheet.cell(row=row, column=column).number_format = "0%"


def _scenario_values(analysis: BugAnalysis) -> list[list[object]]:
    return [
        [
            scenario.test_case_id,
            scenario.category,
            scenario.type,
            scenario.scenario,
            scenario.expected_result,
            scenario.purpose,
            scenario.priority,
        ]
        for scenario in analysis.suggested_test_scenarios
    ]


def _save(workbook: Workbook) -> bytes:
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def single_report(
    bug: dict[str, str], analysis: BugAnalysis, exported_at: datetime
) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Bug Analysis"
    sheet.column_dimensions["A"].width = 24
    sheet.column_dimensions["B"].width = 90

    fields = [
        ("Bug Report", None),
        ("Title", bug.get("title", "")),
        ("Description", bug.get("description", "")),
        ("Steps to Reproduce", _numbered(bug.get("steps_to_reproduce", ""))),
        ("Expected Result", bug.get("expected_result", "")),
        ("Actual Result", bug.get("actual_result", "")),
        ("", None),
        ("Analysis", None),
        ("Severity", analysis.severity),
        ("Priority", analysis.priority),
        ("Category", analysis.category),
        ("Confidence", analysis.confidence),
        ("Impact", analysis.impact),
        ("Possible Root Cause", analysis.possible_root_cause),
        ("Missing Information", _lines(analysis.missing_information) or "None"),
        ("Visual Evidence", analysis.visual_evidence or "No screenshot"),
        ("", None),
        ("Exported", exported_at.strftime("%Y-%m-%d %H:%M")),
    ]
    for row, (label, value) in enumerate(fields, start=1):
        label_cell = _set(sheet, row, 1, label)
        if value is None:
            label_cell.font = SECTION_FONT
            continue
        label_cell.font = Font(bold=True)
        cell = _set(sheet, row, 2, value)
        if label == "Confidence":
            cell.number_format = "0%"
            cell.alignment = Alignment(horizontal="left", vertical="top")

    scenarios = workbook.create_sheet("Test Scenarios")
    _table(scenarios, SCENARIO_COLUMNS, _scenario_values(analysis))

    return _save(workbook)


def batch_report(
    bugs: list[dict],
    source_name: Optional[str],
    exported_at: datetime,
) -> bytes:
    """bugs: the batch result items (row, status, bug, analysis, error,
    duplicate_of), in the order shown on the page."""
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    _batch_summary(summary, bugs, source_name, exported_at)

    sheet = workbook.create_sheet("Bugs")
    is_csv = (source_name or "").lower().endswith(".csv")
    row_title = "File Row" if is_csv else "Excel Row"
    columns = [
        ("Bug", 7),
        (row_title, 11),
        ("Status", 13),
        ("Title", 36),
        ("Description", 45),
        ("Steps to Reproduce", 45),
        ("Expected Result", 32),
        ("Actual Result", 32),
        ("Severity", 11),
        ("Priority", 9),
        ("Category", 20),
        ("Confidence", 11),
        ("Impact", 45),
        ("Possible Root Cause", 45),
        ("Missing Information", 40),
        ("Duplicate Of", 13),
        ("Error", 45),
    ]
    rows = []
    for number, item in enumerate(bugs, start=1):
        values = item["bug"]
        analysis: Optional[BugAnalysis] = item["analysis"]
        rows.append(
            [
                number,
                item["row"],
                STATUS_LABELS[item["status"]],
                values.get("title", ""),
                values.get("description", ""),
                _numbered(values.get("steps_to_reproduce", "")),
                values.get("expected_result", ""),
                values.get("actual_result", ""),
                analysis.severity if analysis else None,
                analysis.priority if analysis else None,
                analysis.category if analysis else None,
                analysis.confidence if analysis else None,
                analysis.impact if analysis else None,
                analysis.possible_root_cause if analysis else None,
                _lines(analysis.missing_information) if analysis else None,
                f"Bug {item['duplicate_of']}" if item["duplicate_of"] else None,
                item["error"],
            ]
        )
    _table(sheet, columns, rows)
    _confidence(sheet, column=12)
    sheet["A1"].comment = Comment(
        "The bug's number on the page (Bug 1, Bug 2...), also used in "
        "Duplicate Of.",
        "AI Bug Analyzer",
    )
    sheet["B1"].comment = Comment(
        f"The row of the bug in the uploaded {'CSV' if is_csv else 'Excel'} "
        "file. Empty and hidden rows are skipped, so the numbers can have gaps.",
        "AI Bug Analyzer",
    )

    scenarios = workbook.create_sheet("Test Scenarios")
    _table(
        scenarios,
        [("Bug", 7), ("Bug Title", 30), *SCENARIO_COLUMNS],
        (
            [number, item["bug"].get("title", ""), *values]
            for number, item in enumerate(bugs, start=1)
            if item["analysis"]
            for values in _scenario_values(item["analysis"])
        ),
    )

    return _save(workbook)


def _batch_summary(
    sheet: Worksheet,
    bugs: list[dict],
    source_name: Optional[str],
    exported_at: datetime,
) -> None:
    sheet.column_dimensions["A"].width = 26
    sheet.column_dimensions["B"].width = 40

    statuses = Counter(item["status"] for item in bugs)
    analyses = [item["analysis"] for item in bugs if item["analysis"]]

    rows: list[tuple[str, object]] = [
        ("Batch Analysis Report", None),
        ("Source file", source_name or "—"),
        ("Exported", exported_at.strftime("%Y-%m-%d %H:%M")),
        ("", None),
        ("Overview", None),
        ("Bugs", len(bugs)),
        *((STATUS_LABELS[status], statuses[status]) for status in STATUS_LABELS),
    ]
    for title, key, order in (
        ("Severity", "severity", ("CRITICAL", "HIGH", "MEDIUM", "LOW")),
        ("Priority", "priority", ("P1", "P2", "P3", "P4")),
        ("Category", "category", ()),
    ):
        counts = Counter(getattr(analysis, key) for analysis in analyses)
        if not counts:
            continue
        rows += [("", None), (title, None)]
        ordered = [value for value in order if value in counts] or [
            value for value, _ in counts.most_common()
        ]
        rows += [(value, counts[value]) for value in ordered]

    for row, (label, value) in enumerate(rows, start=1):
        label_cell = _set(sheet, row, 1, label)
        if value is None:
            label_cell.font = SECTION_FONT if row > 1 else Font(bold=True, size=14)
            continue
        cell = _set(sheet, row, 2, value)
        cell.alignment = Alignment(horizontal="left", vertical="top")
