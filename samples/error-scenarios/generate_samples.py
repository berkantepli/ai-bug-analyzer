"""Generate example files for the errors a user can run into.

Run from the project root:  .venv/bin/python samples/error-scenarios/generate_samples.py
"""

import os
import struct
import zlib
from pathlib import Path

from openpyxl import Workbook

ROOT = Path(__file__).parent

HEADERS = ["Title", "Description", "Steps to Reproduce", "Expected Result", "Actual Result"]
OPTIONAL_HEADERS = ["Severity", "Priority", "Category"]

VALID_BUGS = [
    [
        "Login button does nothing on Safari",
        "Clicking the login button on the sign-in page has no effect in Safari 17.",
        "Open https://app.example.com/login in Safari 17\nEnter a valid email and password\nClick Login",
        "The user is signed in and redirected to the dashboard.",
        "Nothing happens; the console shows TypeError: undefined is not a function.",
    ],
    [
        "Cart total ignores discount code",
        "Applying the SAVE10 code shows a success message but the total is unchanged.",
        "Add any product to the cart\nEnter SAVE10 in the discount field\nClick Apply",
        "The cart total is reduced by 10%.",
        "A success toast appears but the total stays at the original price.",
    ],
    [
        "Profile photo upload fails for PNG files",
        "Uploading a PNG profile photo returns HTTP 500 while JPEG works.",
        "Go to Settings > Profile\nChoose a 2 MB PNG file\nClick Save",
        "The new profile photo is shown.",
        "An error page with HTTP 500 appears and the old photo stays.",
    ],
]

# Longer than the 5000 characters a field may have; also used for
# "__LONG_LOG__" in single/single_scenarios.json.
LONG_LOG = "The app crashes while exporting. Log: " + (
    "java.lang.OutOfMemoryError at ReportExporter.write(ReportExporter.java:120) " * 70
)


def new_sheet(headers=HEADERS):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Bugs"
    sheet.append(headers)
    return workbook, sheet


def save(workbook, batch, name):
    workbook.save(batch / name)


# --- Batch: file-level errors (the whole file is rejected) -----------------


def missing_columns(batch):
    workbook, sheet = new_sheet(["Title", "Description", "Steps to Reproduce"])
    for bug in VALID_BUGS:
        sheet.append(bug[:3])
    save(workbook, batch, "01_eksik_sutunlar.xlsx")


def header_only(batch):
    workbook, _ = new_sheet()
    save(workbook, batch, "02_sadece_baslik_satiri.xlsx")


def all_rows_hidden(batch):
    workbook, sheet = new_sheet()
    for bug in VALID_BUGS:
        sheet.append(bug)
    for row in range(2, 2 + len(VALID_BUGS)):
        sheet.row_dimensions[row].hidden = True
    save(workbook, batch, "03_tum_satirlar_gizli.xlsx")


def too_many_records(batch):
    workbook, sheet = new_sheet()
    for number in range(1, 502):
        sheet.append([
            f"Search result {number} shows the wrong price",
            f"Product number {number} shows a different price in search than on its page.",
            f"Search for product {number}\nCompare the price with the product page",
            "Both pages show the same price.",
            f"The search page shows a price that is {number} cents lower.",
        ])
    save(workbook, batch, "04_500den_fazla_kayit.xlsx")


def too_many_rows(batch):
    workbook, sheet = new_sheet()
    # 2001 rows with data; the row limit is checked before the record limit.
    for number in range(1, 2002):
        sheet.append([f"Note {number}"])
    save(workbook, batch, "05_2000den_fazla_dolu_satir.xlsx")


def too_many_formatted_rows(batch):
    from openpyxl.styles import PatternFill

    workbook, sheet = new_sheet()
    sheet.append(VALID_BUGS[0])
    fill = PatternFill("solid", fgColor="FFFF00")
    # Empty rows that only carry a fill color, far below the data.
    for row in range(3, 50010):
        sheet.cell(row=row, column=1).fill = fill
    save(workbook, batch, "06_50000den_fazla_bicimli_bos_satir.xlsx")


def password_protected(batch):
    # Encrypted Office files are OLE containers; the app recognizes them by
    # their first bytes. This is a stand-in, not a real encrypted file.
    (batch / "07_sifreli_dosya.xlsx").write_bytes(
        b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 4096
    )


def corrupt_file(batch):
    (batch / "08_bozuk_dosya.xlsx").write_bytes(
        b"PK\x03\x04 this is not really a zip archive " * 50
    )


def unsupported_format(batch):
    (batch / "09_desteklenmeyen_format.txt").write_text(
        "Title;Description\nLogin fails;Cannot log in\n", encoding="utf-8"
    )


def empty_file(batch):
    (batch / "10_bos_dosya.xlsx").write_bytes(b"")


def too_large(batch):
    (batch / "11_5mb_ustu_dosya.xlsx").write_bytes(os.urandom(5 * 1024 * 1024 + 1024))


def turkish_headers(batch):
    workbook, sheet = new_sheet(
        ["Başlık", "Açıklama", "Tekrar Adımları", "Beklenen Sonuç", "Gerçekleşen Sonuç"]
    )
    sheet.append([
        "Giriş butonu Safari'de çalışmıyor",
        "Safari 17'de giriş sayfasındaki butona tıklayınca hiçbir şey olmuyor.",
        "Safari 17 ile giriş sayfasını aç\nGeçerli e-posta ve şifre gir\nGiriş'e tıkla",
        "Kullanıcı giriş yapar ve panele yönlendirilir.",
        "Hiçbir şey olmuyor; konsolda TypeError görünüyor.",
    ])
    save(workbook, batch, "12_turkce_basliklar_ollama_gerekir.xlsx")


def unmatchable_headers(batch):
    workbook, sheet = new_sheet(["Kolon A", "Kolon B", "Kolon C", "Kolon D", "Kolon E"])
    sheet.append(["1", "2", "3", "4", "5"])
    sheet.append(["6", "7", "8", "9", "10"])
    save(workbook, batch, "13_anlamsiz_basliklar.xlsx")


# --- Batch: row-level errors (the file is analyzed, bad rows FAILED) -------


ROW_SCENARIOS = [
    # (label, title, description, steps, expected, actual, severity, priority, category)
    ("Geçerli", *VALID_BUGS[0], "Blocker", "Highest", "UI"),
    ("Eksik değer (Expected boş)",
     "Search returns no results for exact names",
     "Searching for an existing product by its exact name returns nothing.",
     "Open the search page\nType 'Blue Kettle'\nPress Enter",
     "", "The page says 'No results found'.", "", "", ""),
    ("Sadece boşluk",
     "Logout link missing on mobile", "   ",
     "Open the app on a phone\nOpen the menu", "A Logout link is shown.",
     "There is no Logout link in the menu.", "", "", ""),
    ("Okunamaz metin",
     "Xq7#vL@ zzkq", "asdkj qweoiu zxcvb", "asdfasdf", "aaaaaa", "qwrtzp mnbvc", "", "", ""),
    ("Placeholder metin",
     "test", "test test", "test", "n/a", "tbd", "", "", ""),
    ("Lorem ipsum",
     "Lorem ipsum dolor", "Lorem ipsum dolor sit amet consectetur.",
     "Lorem\nIpsum", "Dolor sit amet.", "Consectetur adipiscing elit.", "", "", ""),
    ("Aynı metin (Expected = Actual)",
     "Order confirmation email not sent",
     "After placing an order no confirmation email arrives.",
     "Place an order\nCheck the inbox",
     "The confirmation email arrives.", "The confirmation email arrives.", "", "", ""),
    ("5000 karakterden uzun alan",
     "Crash when exporting a large report",
     LONG_LOG,
     "Open Reports\nChoose Export", "The report is exported.", "The app crashes.", "", "", ""),
    ("Tekrar eden kayıt (1. satırın kopyası)", *VALID_BUGS[0], "", "", ""),
    ("Tekrar (büyük/küçük harf ve noktalama farklı)",
     VALID_BUGS[0][0].upper(), VALID_BUGS[0][1].replace(".", "!!"), *VALID_BUGS[0][2:], "", "", ""),
    ("Bug raporu değil (LLM reddeder)",
     "Chocolate cake recipe",
     "Mix flour, sugar, cocoa and eggs, then bake for 35 minutes at 180 degrees.",
     "Preheat the oven\nMix the ingredients\nBake",
     "A soft and moist cake.", "The cake was slightly dry on top.", "", "", ""),
    ("Belirsiz ama analiz edilir (düşük güven)",
     "It does not work", "The page does not work.", "Open the page",
     "It should work.", "It does not work properly.", "", "", ""),
    ("Bilinmeyen Severity/Priority (LLM değeri kalır)",
     *VALID_BUGS[1], "Very Bad", "ASAP", "Checkout"),
    ("Jira değerleri eşlenir (Major→HIGH, Low→P4)",
     *VALID_BUGS[2], "Major", "Low", "Upload"),
]


def row_errors(batch):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Bugs"
    # Notes above the header: the header may be anywhere in the first 20 rows.
    sheet.append(["Hata senaryoları örnek dosyası - başlık satırı 3. satırda"])
    sheet.append([])
    sheet.append(["Scenario (ignored)", *HEADERS, *OPTIONAL_HEADERS])
    for scenario in ROW_SCENARIOS:
        sheet.append(list(scenario))

    # A formula without a saved result (files written by scripts).
    sheet.append([
        "Formül (kaydedilmiş sonuç yok)", "Header overlaps logo on tablets",
        '=CONCATENATE("The header ","overlaps the logo")',
        "Open the home page on an iPad", "The logo is fully visible.",
        "The header covers half of the logo.", "", "", "",
    ])
    # A hidden row is skipped and counted above the results.
    sheet.append([
        "Gizli satır (atlanır)", "Hidden row bug about date picker",
        "The date picker opens behind the dialog.", "Open the booking dialog\nClick the date field",
        "The date picker opens above the dialog.", "It opens behind the dialog.", "", "", "",
    ])
    sheet.row_dimensions[sheet.max_row].hidden = True

    sheet.column_dimensions["A"].width = 42
    for column in "BCDEF":
        sheet.column_dimensions[column].width = 40
    save(workbook, batch, "20_satir_bazli_hatalar.xlsx")


def csv_semicolon_cp1254(batch):
    lines = [";".join(["Senaryo", *HEADERS])]
    rows = [
        ["Geçerli (Türkçe karakterli)", "Ödeme sayfası açılmıyor",
         "Sepetten ödemeye geçince sayfa boş kalıyor.",
         "Sepete ürün ekle / Ödemeye geç", "Ödeme formu görünür.", "Sayfa beyaz ve boş kalıyor."],
        ["Eksik değer", "Şifre sıfırlama e-postası gelmiyor", "", "Şifremi unuttum'a tıkla",
         "E-posta gelir.", "E-posta gelmiyor."],
        ["Placeholder", "deneme", "deneme deneme", "deneme", "test", "asd"],
    ]
    for row in rows:
        lines.append(";".join(row))
    (batch / "21_noktali_virgul_cp1254.csv").write_bytes(
        "\r\n".join(lines).encode("cp1254")
    )


def csv_missing_column(batch):
    (batch / "22_csv_eksik_sutun.csv").write_text(
        "Title,Description,Steps to Reproduce,Expected Result\n"
        "Login fails,Cannot log in with a valid password,Open login page,User logs in\n",
        encoding="utf-8",
    )


# --- Single bug: screenshots -----------------------------------------------


def png(width=4, height=4, color=(220, 40, 40)):
    def chunk(kind, data):
        return (
            struct.pack(">I", len(data)) + kind + data
            + struct.pack(">I", zlib.crc32(kind + data) & 0xFFFFFFFF)
        )

    raw = b"".join(b"\x00" + bytes(color) * width for _ in range(height))
    return (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0))
        + chunk(b"IDAT", zlib.compress(raw))
        + chunk(b"IEND", b"")
    )


def screenshots(directory):
    directory.mkdir(parents=True, exist_ok=True)
    (directory / "gecerli.png").write_bytes(png(64, 64))
    for number in range(1, 7):
        (directory / f"alti_adet_{number}.png").write_bytes(png(8, 8))
    # Starts like a PNG but is cut off: passes the upload check, Ollama fails.
    (directory / "hasarli_yarim.png").write_bytes(png(64, 64)[:40])
    # A text file renamed to .png: rejected by its first bytes.
    (directory / "aslinda_metin.png").write_text("This is not an image.\n")
    # A PNG padded past 10 MB.
    (directory / "10mb_ustu.png").write_bytes(png(8, 8) + b"\x00" * (10 * 1024 * 1024 + 1))


BATCH_BUILDERS = (
    missing_columns, header_only, all_rows_hidden, too_many_records,
    too_many_rows, too_many_formatted_rows, password_protected, corrupt_file,
    unsupported_format, empty_file, too_large, turkish_headers,
    unmatchable_headers, row_errors, csv_semicolon_cp1254, csv_missing_column,
)


def generate(root=ROOT):
    """Write batch/ and single/screenshots/ under root; the tests use a
    temporary root."""
    batch = root / "batch"
    batch.mkdir(parents=True, exist_ok=True)
    for build in BATCH_BUILDERS:
        build(batch)
    screenshots(root / "single" / "screenshots")
    return batch, root / "single" / "screenshots"


def main():
    batch, shots = generate()
    print("Generated", len(list(batch.iterdir())), "batch files and",
          len(list(shots.iterdir())), "screenshots")


if __name__ == "__main__":
    main()
