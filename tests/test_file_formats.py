from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook
from openpyxl.styles import Border, Side

from app.main import app
from helpers import FAKE_ANALYSIS, HEADER


client = TestClient(app)

LOGIN = ("Login fails", "Cannot sign in", "1. Login", "Dashboard", "Error")
LOGOUT = ("Logout fails", "Cannot sign out", "1. Logout", "Login page", "Error")


@pytest.fixture(autouse=True)
def fake_llm(monkeypatch) -> list:
    titles = []

    async def fake_analyze_with_llm(bug, screenshots=None):
        titles.append(bug.title)
        return FAKE_ANALYSIS

    monkeypatch.setattr("app.api.bugs.analyze_with_llm", fake_analyze_with_llm)
    return titles


def post(name: str, content: bytes):
    return client.post("/bugs/batch", files={"file": (name, content)})


@pytest.mark.parametrize(
    ("content", "encoding"),
    [
        ("Bug Title,Description,Steps to Reproduce,Expected Result,Actual Result\n"
         'Login fails,Cannot sign in,"1. Login\n2. Submit",Dashboard,Error\n', "utf-8"),
        ("Bug Title;Description;Steps to Reproduce;Expected Result;Actual Result\n"
         "Giriş başarısız;Giriş yapılamıyor;1. Giriş;Panel açılır;Hata\n", "cp1254"),
        ("﻿Bug Title\tDescription\tSteps to Reproduce\tExpected Result\tActual Result\n"
         "Login fails\tCannot sign in\t1. Login\tDashboard\tError\n", "utf-8"),
    ],
    ids=["comma", "semicolon in Turkish Windows encoding", "tab with BOM"],
)
def test_csv_files_are_read(content, encoding) -> None:
    response = post("bugs.csv", content.encode(encoding))

    assert response.status_code == 200
    bug = response.json()["bugs"][0]
    assert bug["status"] == "analyzed"
    assert bug["bug"]["title"] in {"Login fails", "Giriş başarısız"}


def test_empty_file_is_reported() -> None:
    response = post("bugs.xlsx", b"")

    assert response.status_code == 422
    assert response.json()["detail"] == "The uploaded file is empty."


def test_password_protected_xlsx_is_reported() -> None:
    # Encrypted .xlsx files are OLE containers instead of zip files.
    content = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 512

    response = post("bugs.xlsx", content)

    assert response.status_code == 422
    assert "password-protected" in response.json()["detail"]


def test_hidden_rows_are_skipped(fake_llm) -> None:
    workbook = Workbook()
    sheet = workbook.active
    for row in (HEADER, LOGIN, LOGOUT):
        sheet.append(row)
    sheet.row_dimensions[2].hidden = True
    output = BytesIO()
    workbook.save(output)

    response = post("bugs.xlsx", output.getvalue())

    data = response.json()
    assert fake_llm == ["Logout fails"]
    assert [bug["row"] for bug in data["bugs"]] == [3]
    assert data["skipped_hidden_rows"] == 1


def test_only_hidden_rows_explains_why_nothing_is_analyzed() -> None:
    workbook = Workbook()
    sheet = workbook.active
    for row in (HEADER, LOGIN):
        sheet.append(row)
    sheet.row_dimensions[2].hidden = True
    output = BytesIO()
    workbook.save(output)

    response = post("bugs.xlsx", output.getvalue())

    assert response.status_code == 422
    assert response.json()["detail"] == (
        "The Excel file contains no bug records. "
        "1 hidden rows were skipped; unhide them to analyze them."
    )


def test_csv_messages_name_the_csv_file() -> None:
    response = post("bugs.csv", b"Name,Notes\nx,y\n")

    assert response.status_code == 422
    assert response.json()["detail"].startswith("Missing required CSV columns:")


def test_empty_formatted_rows_do_not_count_towards_the_row_limit(monkeypatch) -> None:
    monkeypatch.setattr("app.services.batch_analyzer.MAX_SHEET_ROWS", 3)
    workbook = Workbook()
    sheet = workbook.active
    for row in (HEADER, LOGIN, LOGOUT):
        sheet.append(row)
    # Borders drawn far below the data, as in a formatted table.
    for row_number in range(4, 40):
        sheet.cell(row=row_number, column=1).border = Border(top=Side(style="thin"))
    output = BytesIO()
    workbook.save(output)

    response = post("bugs.xlsx", output.getvalue())

    assert response.status_code == 200
    assert [bug["row"] for bug in response.json()["bugs"]] == [2, 3]


def test_csv_blank_lines_do_not_count_towards_the_row_limit(monkeypatch) -> None:
    monkeypatch.setattr("app.services.batch_analyzer.MAX_SHEET_ROWS", 3)
    content = ",".join(HEADER) + "\n" + ",".join(LOGIN) + "\n" + "\n" * 20

    response = post("bugs.csv", content.encode())

    assert response.status_code == 200


@pytest.mark.parametrize(
    "content",
    [
        ("sep=;\n" + ";".join(HEADER) + "\n" + ";".join(LOGIN) + "\n").encode(),
        ("\t".join(HEADER) + "\n" + "\t".join(LOGIN) + "\n").encode("utf-16"),
    ],
    ids=["Excel separator line", "UTF-16"],
)
def test_csv_exports_from_excel_are_read(content) -> None:
    response = post("bugs.csv", content)

    assert response.status_code == 200
    assert response.json()["bugs"][0]["bug"]["title"] == "Login fails"