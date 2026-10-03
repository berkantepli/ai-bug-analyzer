import asyncio
from io import BytesIO

from openpyxl import Workbook
from fastapi.testclient import TestClient
import httpx

from app.api import bugs
from app.main import app
from app.schemas.analysis import BugAnalysis
from app.services.llm_analyzer import InvalidBugReportError, OllamaUnavailableError


client = TestClient(app)


FAKE_ANALYSIS = BugAnalysis(
    severity="HIGH",
    priority="P2",
    category="Authentication",
    impact="Users cannot sign in.",
    possible_root_cause="The login request is rejected.",
    suggested_test_scenarios=[],
    missing_information=[],
    confidence=0.8,
)


def excel_bytes(rows: list[tuple[object, ...]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    for row in rows:
        sheet.append(row)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_batch_endpoint_reads_bug_reports_from_xlsx(monkeypatch) -> None:
    async def fake_analyze_with_llm(bug, screenshots=None):
        return FAKE_ANALYSIS

    monkeypatch.setattr("app.api.bugs.analyze_with_llm", fake_analyze_with_llm)

    content = excel_bytes(
        [
            (
                "Bug Title",
                "Description",
                "Steps to Reproduce",
                "Expected Result",
                "Actual Result",
            ),
            (
                "Login fails",
                "Cannot sign in",
                "1. Open login\n2. Submit",
                "Dashboard opens",
                "Error appears",
            ),
        ]
    )

    response = client.post("/bugs/batch", files={"file": ("bugs.xlsx", content)})

    assert response.status_code == 200
    data = response.json()
    assert data["total"] == 1
    assert data["completed"] == 1
    assert data["failed"] == 0
    assert data["bugs"] == [
        {
            "row": 2,
            "status": "analyzed",
            "title": "Login fails",
            "bug": {
                "title": "Login fails",
                "description": "Cannot sign in",
                "steps_to_reproduce": "1. Open login\n2. Submit",
                "expected_result": "Dashboard opens",
                "actual_result": "Error appears",
            },
            "analysis": FAKE_ANALYSIS.model_dump(),
            "error": None,
            "duplicate_of": None,
        }
    ]


def test_batch_endpoint_reports_missing_columns() -> None:
    content = excel_bytes([("Title", "Description"), ("Broken", "Details")])

    response = client.post("/bugs/batch", files={"file": ("bugs.xlsx", content)})

    assert response.status_code == 422
    assert "Missing required Excel columns" in response.json()["detail"]


def test_batch_endpoint_rejects_unsupported_file_types() -> None:
    response = client.post(
        "/bugs/batch", files={"file": ("bugs.csv", b"title,description")}
    )

    assert response.status_code == 415
    assert ".xlsx and .xls" in response.json()["detail"]


HEADER = (
    "Bug Title",
    "Description",
    "Steps to Reproduce",
    "Expected Result",
    "Actual Result",
)


def test_batch_endpoint_marks_rows_with_missing_values_as_failed(monkeypatch) -> None:
    analyzed_titles = []

    async def fake_analyze_with_llm(bug, screenshots=None):
        analyzed_titles.append(bug.title)
        return FAKE_ANALYSIS

    monkeypatch.setattr("app.api.bugs.analyze_with_llm", fake_analyze_with_llm)

    content = excel_bytes(
        [
            HEADER,
            ("Login fails", "", "1. Open login", "Dashboard opens", "Error appears"),
            ("Logout fails", "Cannot sign out", "1. Logout", "Login page", "Error"),
        ]
    )

    response = client.post("/bugs/batch", files={"file": ("bugs.xlsx", content)})

    assert response.status_code == 200
    data = response.json()
    assert data["completed"] == 1
    assert data["failed"] == 1
    assert analyzed_titles == ["Logout fails"]

    failed_bug, analyzed_bug = data["bugs"]
    assert failed_bug["row"] == 2
    assert failed_bug["status"] == "failed"
    assert failed_bug["analysis"] is None
    assert failed_bug["bug"]["title"] == "Login fails"
    assert failed_bug["error"] == "Missing required values: description."
    assert analyzed_bug["status"] == "analyzed"


def test_batch_endpoint_marks_llm_errors_as_failed(monkeypatch) -> None:
    async def fake_analyze_with_llm(bug, screenshots=None):
        if bug.title == "Login fails":
            raise RuntimeError("The LLM returned invalid JSON.")
        return FAKE_ANALYSIS

    monkeypatch.setattr("app.api.bugs.analyze_with_llm", fake_analyze_with_llm)

    content = excel_bytes(
        [
            HEADER,
            ("Login fails", "Cannot sign in", "1. Login", "Dashboard", "Error"),
            ("Logout fails", "Cannot sign out", "1. Logout", "Login page", "Error"),
        ]
    )

    response = client.post("/bugs/batch", files={"file": ("bugs.xlsx", content)})

    assert response.status_code == 200
    data = response.json()
    assert data["completed"] == 1
    assert data["failed"] == 1

    failed_bug, analyzed_bug = data["bugs"]
    assert failed_bug["status"] == "failed"
    assert failed_bug["analysis"] is None
    assert failed_bug["error"] == (
        "The LLM could not analyze this bug: The LLM returned invalid JSON."
    )
    assert analyzed_bug["status"] == "analyzed"
    assert analyzed_bug["analysis"] == FAKE_ANALYSIS.model_dump()


def test_batch_endpoint_marks_unreadable_rows_as_failed(monkeypatch) -> None:
    analyzed_titles = []

    async def fake_analyze_with_llm(bug, screenshots=None):
        analyzed_titles.append(bug.title)
        return FAKE_ANALYSIS

    monkeypatch.setattr("app.api.bugs.analyze_with_llm", fake_analyze_with_llm)

    content = excel_bytes(
        [
            HEADER,
            (
                "Xq7#vL@ zR9$ kT~pW 4&mN",
                "Fj%8 qwZ!x 9Lp#r Tz@4k Vb^7n",
                "1. Wq#8z Kp@3 lX!",
                "Lx~6p Qr#4 zM@ Vn$1",
                "Dz*2q Pj+9 xK? Ym=5",
            ),
            (
                "TypeError on profile page",
                "TypeError: Cannot read property 'x' of undefined at app.js:42",
                "1. Open /profile",
                "Profile page loads",
                "API returns 500 on GET /api/v1/users?id=12",
            ),
        ]
    )

    response = client.post("/bugs/batch", files={"file": ("bugs.xlsx", content)})

    assert response.status_code == 200
    data = response.json()
    assert data["completed"] == 1
    assert data["failed"] == 1
    assert analyzed_titles == ["TypeError on profile page"]

    unreadable_bug = data["bugs"][0]
    assert unreadable_bug["status"] == "failed"
    assert unreadable_bug["error"] == "The bug report text is unreadable."


def test_batch_endpoint_marks_placeholder_and_invalid_reports_as_failed(
    monkeypatch,
) -> None:
    async def fake_analyze_with_llm(bug, screenshots=None):
        if bug.title == "Chocolate cake":
            raise InvalidBugReportError("The text is a cake recipe.")
        return FAKE_ANALYSIS

    monkeypatch.setattr("app.api.bugs.analyze_with_llm", fake_analyze_with_llm)

    content = excel_bytes(
        [
            HEADER,
            ("test", "test test", "1. test", "test", "test"),
            ("Chocolate cake", "Mix flour and sugar", "1. Bake", "Moist", "Tasty"),
            ("Logout fails", "Cannot sign out", "1. Logout", "Login page", "Error"),
        ]
    )

    response = client.post("/bugs/batch", files={"file": ("bugs.xlsx", content)})

    assert response.status_code == 200
    placeholder, recipe, valid = response.json()["bugs"]
    assert placeholder["error"] == "The bug report contains only placeholder text."
    assert recipe["status"] == "failed"
    assert recipe["error"] == "Not a valid bug report: The text is a cake recipe."
    assert valid["status"] == "analyzed"


def test_batch_endpoint_marks_duplicate_bugs(monkeypatch) -> None:
    analyzed_titles = []

    async def fake_analyze_with_llm(bug, screenshots=None):
        analyzed_titles.append(bug.title)
        return FAKE_ANALYSIS

    monkeypatch.setattr("app.api.bugs.analyze_with_llm", fake_analyze_with_llm)

    login = ("Login fails", "Cannot sign in", "1. Login", "Dashboard", "Error")
    logout = ("Logout fails", "Cannot sign out", "1. Logout", "Login page", "Error")
    login_again = ("LOGIN FAILS", "Cannot  sign in.", "1. login", "Dashboard!", "error")

    content = excel_bytes([HEADER, login, logout, login_again, logout])

    response = client.post("/bugs/batch", files={"file": ("bugs.xlsx", content)})

    assert response.status_code == 200
    data = response.json()
    assert analyzed_titles == ["Login fails", "Logout fails"]
    assert data["completed"] == 2
    assert data["failed"] == 0
    assert data["duplicates"] == 2
    assert [bug["status"] for bug in data["bugs"]] == [
        "analyzed",
        "analyzed",
        "duplicate",
        "duplicate",
    ]
    assert [bug["duplicate_of"] for bug in data["bugs"]] == [None, None, 1, 2]
    assert data["bugs"][2]["analysis"] is None


def test_batch_endpoint_marks_rows_with_identical_fields_as_failed(
    monkeypatch,
) -> None:
    async def fake_analyze_with_llm(bug, screenshots=None):
        return FAKE_ANALYSIS

    monkeypatch.setattr("app.api.bugs.analyze_with_llm", fake_analyze_with_llm)

    content = excel_bytes(
        [HEADER, ("Login fails", "Cannot sign in", "1. Login", "Error", "Error")]
    )

    response = client.post("/bugs/batch", files={"file": ("bugs.xlsx", content)})

    bug = response.json()["bugs"][0]
    assert bug["status"] == "failed"
    assert bug["error"] == "Expected result and actual result contain the same text."


def test_batch_endpoint_marks_too_long_rows_as_failed(monkeypatch) -> None:
    async def fake_analyze_with_llm(bug, screenshots=None):
        return FAKE_ANALYSIS

    monkeypatch.setattr("app.api.bugs.analyze_with_llm", fake_analyze_with_llm)

    content = excel_bytes(
        [
            HEADER,
            ("Login fails", "Long text. " * 1000, "1. Login", "Dashboard", "Error"),
        ]
    )

    response = client.post("/bugs/batch", files={"file": ("bugs.xlsx", content)})

    bug = response.json()["bugs"][0]
    assert bug["status"] == "failed"
    assert bug["error"] == "Description is longer than 5000 characters."


def test_batch_stops_sending_bugs_once_ollama_is_unavailable(monkeypatch) -> None:
    analyzed_titles = []

    async def fake_analyze_with_llm(bug, screenshots=None):
        analyzed_titles.append(bug.title)
        raise OllamaUnavailableError("Could not connect to Ollama.")

    monkeypatch.setattr("app.api.bugs.analyze_with_llm", fake_analyze_with_llm)

    content = excel_bytes(
        [
            HEADER,
            ("Login fails", "Cannot sign in", "1. Login", "Dashboard", "Error"),
            ("Logout fails", "Cannot sign out", "1. Logout", "Login page", "Error"),
            ("Search fails", "No results", "1. Search", "Results", "Empty list"),
        ]
    )

    response = client.post("/bugs/batch", files={"file": ("bugs.xlsx", content)})

    data = response.json()
    assert analyzed_titles == ["Login fails"]
    assert data["failed"] == 0
    assert data["not_analyzed"] == 3
    assert data["stopped_reason"] == (
        "Ollama became unavailable during the batch: Could not connect to Ollama."
    )
    # The bug that hit the error and the skipped ones look the same; the
    # Ollama error itself is reported once in stopped_reason.
    assert [bug["status"] for bug in data["bugs"]] == ["not_analyzed"] * 3
    assert [bug["error"] for bug in data["bugs"]] == [bugs.NOT_ANALYZED_ERROR] * 3


def test_batch_without_ollama_errors_has_no_stopped_reason(monkeypatch) -> None:
    async def fake_analyze_with_llm(bug, screenshots=None):
        return FAKE_ANALYSIS

    monkeypatch.setattr("app.api.bugs.analyze_with_llm", fake_analyze_with_llm)

    content = excel_bytes(
        [HEADER, ("Login fails", "Cannot sign in", "1. Login", "Dashboard", "Error")]
    )

    response = client.post("/bugs/batch", files={"file": ("bugs.xlsx", content)})

    assert response.json()["stopped_reason"] is None


def test_batch_endpoint_finds_header_below_report_title(monkeypatch) -> None:
    async def fake_analyze_with_llm(bug, screenshots=None):
        return FAKE_ANALYSIS

    monkeypatch.setattr("app.api.bugs.analyze_with_llm", fake_analyze_with_llm)

    content = excel_bytes(
        [
            ("Sprint 12 Bug Report",),
            ("Exported on 2026-10-02",),
            (),
            HEADER,
            ("Login fails", "Cannot sign in", "1. Login", "Dashboard", "Error"),
            ("Logout fails", "", "1. Logout", "Login page", "Error"),
        ]
    )

    response = client.post("/bugs/batch", files={"file": ("bugs.xlsx", content)})

    assert response.status_code == 200
    data = response.json()
    assert [bug["row"] for bug in data["bugs"]] == [5, 6]
    assert [bug["status"] for bug in data["bugs"]] == ["analyzed", "failed"]


def test_batch_endpoint_keeps_excel_row_numbers_when_sheet_starts_lower(
    monkeypatch,
) -> None:
    async def fake_analyze_with_llm(bug, screenshots=None):
        return FAKE_ANALYSIS

    monkeypatch.setattr("app.api.bugs.analyze_with_llm", fake_analyze_with_llm)

    workbook = Workbook()
    sheet = workbook.active
    for column, value in enumerate(HEADER, start=1):
        sheet.cell(row=3, column=column, value=value)
    for column, value in enumerate(
        ("Login fails", "Cannot sign in", "1. Login", "Dashboard", "Error"),
        start=1,
    ):
        sheet.cell(row=4, column=column, value=value)
    output = BytesIO()
    workbook.save(output)

    response = client.post(
        "/bugs/batch", files={"file": ("bugs.xlsx", output.getvalue())}
    )

    assert response.status_code == 200
    assert response.json()["bugs"][0]["row"] == 4


def test_batch_progress_reports_rejected_rows_up_front(monkeypatch) -> None:
    progress_during_analysis = []

    async def fake_analyze_with_llm(bug, screenshots=None):
        progress_during_analysis.append(dict(bugs.batch_progress["tab-a"]))
        return FAKE_ANALYSIS

    monkeypatch.setattr("app.api.bugs.analyze_with_llm", fake_analyze_with_llm)

    content = excel_bytes(
        [
            HEADER,
            ("Login fails", "", "1. Login", "Dashboard", "Error"),
            ("Logout fails", "Cannot sign out", "1. Logout", "Login page", "Error"),
        ]
    )

    response = client.post(
        "/bugs/batch",
        files={"file": ("bugs.xlsx", content)},
        data={"batch_id": "tab-a"},
    )

    assert response.status_code == 200
    data = response.json()
    assert data["batch_id"] == "tab-a"
    assert data["rejected"] == 1
    assert progress_during_analysis == [
        {
            "status": "processing",
            "total": 2,
            "completed": 0,
            "failed": 1,
            "rejected": 1,
            "duplicates": 0,
            "not_analyzed": 0,
        }
    ]
    assert "tab-a" not in bugs.batch_progress


def test_batch_progress_returns_404_for_unknown_batch() -> None:
    response = client.get("/bugs/batch/unknown/progress")

    assert response.status_code == 404


def test_batch_endpoint_rejects_running_batch_id() -> None:
    bugs.batch_progress["tab-a"] = {"status": "processing"}
    try:
        response = client.post(
            "/bugs/batch",
            files={"file": ("bugs.xlsx", excel_bytes([HEADER]))},
            data={"batch_id": "tab-a"},
        )
    finally:
        bugs.batch_progress.pop("tab-a")

    assert response.status_code == 409


def test_concurrent_batches_keep_separate_progress(monkeypatch) -> None:
    snapshots = {}

    async def fake_analyze_with_llm(bug, screenshots=None):
        # Let the other batch run before this one records its progress.
        await asyncio.sleep(0.01)
        snapshots[bug.title] = {
            batch_id: dict(progress)
            for batch_id, progress in bugs.batch_progress.items()
        }
        return FAKE_ANALYSIS

    monkeypatch.setattr("app.api.bugs.analyze_with_llm", fake_analyze_with_llm)

    tab_a = excel_bytes(
        [HEADER, ("Login fails", "Cannot sign in", "1. Login", "Dashboard", "Error")]
    )
    tab_b = excel_bytes(
        [
            HEADER,
            ("Logout fails", "", "1. Logout", "Login page", "Error"),
            ("Search fails", "No results", "1. Search", "Results", "Empty"),
            ("Upload fails", "Upload error", "1. Upload", "Uploaded", "Error"),
        ]
    )

    async def run_batches():
        transport = httpx.ASGITransport(app=app)
        async with httpx.AsyncClient(
            transport=transport, base_url="http://test"
        ) as async_client:
            return await asyncio.gather(
                async_client.post(
                    "/bugs/batch",
                    files={"file": ("a.xlsx", tab_a)},
                    data={"batch_id": "tab-a"},
                ),
                async_client.post(
                    "/bugs/batch",
                    files={"file": ("b.xlsx", tab_b)},
                    data={"batch_id": "tab-b"},
                ),
            )

    response_a, response_b = asyncio.run(run_batches())

    assert response_a.json()["total"] == 1
    assert response_b.json()["total"] == 3
    assert response_b.json()["failed"] == 1

    progress_while_both_ran = snapshots["Login fails"]
    assert progress_while_both_ran["tab-a"]["total"] == 1
    assert progress_while_both_ran["tab-a"]["failed"] == 0
    assert progress_while_both_ran["tab-b"]["total"] == 3
    assert progress_while_both_ran["tab-b"]["failed"] == 1
    assert bugs.batch_progress == {}
