import zipfile
from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import app
from helpers import FAKE_ANALYSIS


client = TestClient(app)

ANALYSIS = FAKE_ANALYSIS.model_dump(mode="json") | {
    "suggested_test_scenarios": [
        {
            "test_case_id": "TC-01",
            "category": "Functional Validation",
            "type": "Positive",
            "scenario": "Sign in with valid credentials",
            "expected_result": "The dashboard opens",
            "purpose": "Covers the reported flow",
            "priority": "High",
        }
    ],
    "missing_information": ["Browser version"],
}

BUG = {
    "title": "Login fails",
    "description": "Valid users cannot sign in.",
    "steps_to_reproduce": "1. Open the login page\n• Upload a 1.5 GB file",
    "expected_result": "The dashboard opens",
    "actual_result": "An error appears",
}


def workbook(response):
    return load_workbook(BytesIO(response.content))


def test_single_report_contains_the_bug_and_its_analysis() -> None:
    response = client.post(
        "/bugs/export/single", json={"bug": BUG, "analysis": ANALYSIS}
    )

    assert response.status_code == 200
    assert response.headers["content-disposition"].startswith(
        'attachment; filename="bug-analysis-'
    )
    book = workbook(response)
    values = dict(book["Bug Analysis"].iter_rows(values_only=True))
    assert values["Title"] == "Login fails"
    assert values["Steps to Reproduce"] == (
        "1. Open the login page\n2. Upload a 1.5 GB file"
    )
    assert values["Severity"] == "HIGH"
    assert values["Missing Information"] == "• Browser version"
    scenarios = list(book["Test Scenarios"].iter_rows(values_only=True))
    assert scenarios[1][:3] == ("TC-01", "Functional Validation", "Positive")


def test_batch_report_lists_every_bug_with_its_status() -> None:
    bugs = [
        {"row": 2, "status": "analyzed", "bug": BUG, "analysis": ANALYSIS},
        {"row": 3, "status": "duplicate", "bug": BUG, "duplicate_of": 1},
        {
            "row": 4,
            "status": "failed",
            "bug": {"title": "Bad row"},
            "error": "Missing required values: description.",
        },
        {
            "row": 5,
            "status": "not_analyzed",
            "bug": {"title": "Later"},
            "error": "Ollama stopped.",
        },
    ]

    response = client.post(
        "/bugs/export/batch", json={"source_name": "Hata Listesi.xlsx", "bugs": bugs}
    )

    assert response.status_code == 200
    disposition = response.headers["content-disposition"]
    assert 'filename="Hata-Listesi-analysis-' in disposition
    book = workbook(response)
    assert book.sheetnames == ["Summary", "Bugs", "Test Scenarios"]
    summary = dict(row for row in book["Summary"].iter_rows(values_only=True) if row[0])
    assert (summary["Bugs"], summary["Analyzed"], summary["Failed"]) == (4, 1, 1)
    assert (summary["Duplicate"], summary["Not analyzed"]) == (1, 1)
    rows = list(book["Bugs"].iter_rows(values_only=True))
    assert rows[0][:2] == ("Bug", "Excel Row")
    assert "uploaded Excel file" in book["Bugs"]["B1"].comment.text
    assert [row[2] for row in rows[1:]] == [
        "Analyzed",
        "Duplicate",
        "Failed",
        "Not analyzed",
    ]
    assert rows[2][15] == "Bug 1"
    assert rows[3][16] == "Missing required values: description."
    assert len(list(book["Test Scenarios"].iter_rows())) == 2


def test_report_text_never_becomes_a_formula() -> None:
    # Text from users and the LLM must not run as a formula in Excel, and
    # control characters Excel rejects must not break the file.
    bug = BUG | {"title": '=HYPERLINK("http://example.com","Click")'}
    analysis = ANALYSIS | {"impact": "Colors\x1b[31m in a log"}

    response = client.post(
        "/bugs/export/single", json={"bug": bug, "analysis": analysis}
    )

    sheet = zipfile.ZipFile(BytesIO(response.content)).read("xl/worksheets/sheet1.xml")
    assert b"<f>" not in sheet
    values = dict(workbook(response)["Bug Analysis"].iter_rows(values_only=True))
    assert values["Title"] == '=HYPERLINK("http://example.com","Click")'
    assert values["Impact"] == "Colors[31m in a log"


def test_report_requests_are_validated() -> None:
    unknown_status = {"bugs": [{"row": 2, "status": "done", "bug": {}}]}
    unknown_field = {"bug": {"script": "x"}, "analysis": ANALYSIS}

    assert client.post("/bugs/export/batch", json=unknown_status).status_code == 422
    assert client.post("/bugs/export/single", json=unknown_field).status_code == 422


def test_csv_batch_report_names_the_file_row() -> None:
    bugs = [{"row": 2, "status": "analyzed", "bug": BUG, "analysis": ANALYSIS}]

    response = client.post(
        "/bugs/export/batch", json={"source_name": "bugs.csv", "bugs": bugs}
    )

    assert workbook(response)["Bugs"]["B1"].value == "File Row"
